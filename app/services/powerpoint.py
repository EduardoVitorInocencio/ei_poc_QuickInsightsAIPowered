"""Geracao de PowerPoint tabular para insights de vendas.

Este servico nao consulta dados nem decide quando exportar. Ele recebe
``result_id`` validos e cuida somente da composicao visual e persistencia.
"""

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from app.services.query_results import QueryResultStore


class PowerPointService:
    """Acumula um slide tabular por insight em uma apresentacao."""

    def __init__(
        self,
        output_path: str | Path,
        result_store: QueryResultStore,
        max_table_rows: int = 15,
        max_table_columns: int = 8,
    ) -> None:
        """Configura o arquivo e os limites de legibilidade.

        Args:
            output_path: Caminho do arquivo ``.pptx``.
            result_store: Fonte dos DataFrames referenciados pelos insights.
            max_table_rows: Maximo de linhas exibidas por slide.
            max_table_columns: Maximo de colunas aceitas por slide.
        """
        self._base_output_path = Path(output_path)
        self._output_path = self._base_output_path
        self._result_store = result_store
        self._max_table_rows = max_table_rows
        self._max_table_columns = max_table_columns
        self._presentation: Presentation | None = None

    @property
    def output_path(self) -> Path:
        """Retorna o caminho configurado sem criar a apresentacao."""
        return self._output_path

    @property
    def slide_count(self) -> int:
        """Retorna zero ou a quantidade de slides acumulados."""
        if self._presentation is None:
            return 0
        return len(self._presentation.slides)

    def start_report(self) -> None:
        """Descarta a apresentacao anterior e inicia um arquivo em memoria."""
        # O caminho final recebe timestamp para nao sobrescrever exportacoes
        # anteriores e para facilitar auditoria manual.
        self._output_path = self._timestamped_output_path(datetime.now())
        self._presentation = Presentation()
        self._presentation.slide_width = Inches(13.333)
        self._presentation.slide_height = Inches(7.5)

    def create_table_slide(
        self,
        *,
        result_id: str,
        titulo: str,
        headline: str,
    ) -> dict[str, Any]:
        """Adiciona e salva um slide vinculado a um resultado SQL.

        Args:
            result_id: Referencia ao DataFrame do insight.
            titulo: Titulo executivo exibido no topo.
            headline: Principal conclusao exibida antes da tabela.

        Returns:
            Status, caminho, numero do slide e metadados do conteudo.

        Raises:
            ValueError: Se textos, resultado ou dimensoes da tabela forem
                invalidos.
        """
        self._validate_text(titulo, headline)
        # O `result_id` liga o insight ao resultado SQL salvo anteriormente;
        # isso garante que a tabela apresentada é a mesma analisada pelo agente.
        dataframe = self._prepare_table_data(result_id)
        presentation = self._get_presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        self._add_textbox(
            slide=slide,
            text=titulo,
            top=0.3,
            height=0.6,
            font_size=27,
            color=RGBColor(20, 20, 20),
            bold=True,
        )
        self._add_textbox(
            slide=slide,
            text=headline,
            top=0.95,
            height=0.8,
            font_size=16,
            color=RGBColor(55, 55, 55),
        )
        self._add_table(slide=slide, dataframe=dataframe)
        self._save()

        slide_data = {
            "titulo": titulo,
            "headline": headline,
            "columns": list(dataframe.columns),
            "displayed_rows": int(len(dataframe)),
        }
        return {
            "ok": True,
            "pptx_path": str(self._output_path),
            "slide_number": self.slide_count,
            "slide": slide_data,
        }

    def _get_presentation(self) -> Presentation:
        """Retorna a apresentacao atual, criando uma quando necessario."""
        if self._presentation is None:
            self.start_report()
        assert self._presentation is not None
        return self._presentation

    def _prepare_table_data(self, result_id: str) -> pd.DataFrame:
        """Recupera e limita os dados para caberem em um slide."""
        dataframe = self._result_store.get(result_id)
        if dataframe.empty:
            raise ValueError("O resultado SQL esta vazio e nao pode gerar tabela.")
        if len(dataframe.columns) > self._max_table_columns:
            # O limite de colunas existe porque o slide precisa permanecer
            # legível sem exigir rolagem ou redução extrema de fonte.
            raise ValueError(
                f"O resultado possui {len(dataframe.columns)} colunas. "
                f"O limite para o slide e {self._max_table_columns}."
            )
        return dataframe.head(self._max_table_rows).copy()

    @staticmethod
    def _add_table(*, slide: Any, dataframe: pd.DataFrame) -> None:
        """Insere uma tabela nativa do PowerPoint no slide."""
        rows = len(dataframe) + 1
        columns = len(dataframe.columns)
        table = slide.shapes.add_table(
            rows,
            columns,
            Inches(0.7),
            Inches(1.95),
            Inches(11.9),
            Inches(4.9),
        ).table

        # Todas as colunas recebem a mesma largura para manter previsibilidade.
        column_width = Inches(11.9 / columns)
        for column_index in range(columns):
            table.columns[column_index].width = column_width

        for column_index, column_name in enumerate(dataframe.columns):
            cell = table.cell(0, column_index)
            cell.text = str(column_name)
            PowerPointService._format_cell(
                cell,
                font_size=11,
                bold=True,
                font_color=RGBColor(255, 255, 255),
                fill_color=RGBColor(31, 78, 121),
            )

        for row_index, row in enumerate(
            dataframe.itertuples(index=False, name=None),
            start=1,
        ):
            # Linhas alternadas melhoram leitura e ajudam a distinguir células
            # em tabelas densas.
            fill_color = (
                RGBColor(242, 246, 250)
                if row_index % 2 == 0
                else RGBColor(255, 255, 255)
            )
            for column_index, value in enumerate(row):
                cell = table.cell(row_index, column_index)
                cell.text = PowerPointService._format_value(value)
                PowerPointService._format_cell(
                    cell,
                    font_size=10,
                    fill_color=fill_color,
                )

    @staticmethod
    def _format_value(value: Any) -> str:
        """Converte valores do Pandas em texto adequado para celulas."""
        if pd.isna(value):
            return ""
        if isinstance(value, float):
            return f"{value:,.2f}"
        return str(value)

    @staticmethod
    def _format_cell(
        cell: Any,
        *,
        font_size: int,
        fill_color: RGBColor,
        bold: bool = False,
        font_color: RGBColor = RGBColor(35, 35, 35),
    ) -> None:
        """Aplica preenchimento, margens, alinhamento e fonte a uma celula."""
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill_color
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = font_color

    def _save(self) -> None:
        """Cria o diretorio de destino e persiste o estado atual."""
        self._output_path.parent.mkdir(parents=True, exist_ok=True)
        self._get_presentation().save(self._output_path)

    @staticmethod
    def _validate_text(titulo: str, headline: str) -> None:
        """Rejeita slides sem os dois elementos narrativos obrigatorios."""
        if not titulo.strip():
            raise ValueError("O titulo do slide nao pode estar vazio.")
        if not headline.strip():
            raise ValueError("A headline do slide nao pode estar vazia.")

    @staticmethod
    def _add_textbox(
        *,
        slide: Any,
        text: str,
        top: float,
        height: float,
        font_size: int,
        color: RGBColor,
        bold: bool = False,
    ) -> None:
        """Insere uma caixa de texto com tipografia configuravel."""
        textbox = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(top),
            Inches(11.7),
            Inches(height),
        )
        text_frame = textbox.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color

    def _timestamped_output_path(self, generated_at: datetime) -> Path:
        """Cria um nome de arquivo com data e hora para cada exportacao."""
        suffix = self._base_output_path.suffix or ".pptx"
        stem = self._base_output_path.stem
        timestamp = generated_at.strftime("%Y%m%d_%H%M%S")
        # Usamos sempre o nome base para que cada exportacao gere um arquivo
        # novo, sem carregar sufixos acumulados.
        self._output_path = self._base_output_path.with_name(
            f"{stem}_{timestamp}{suffix}"
        )
        return self._output_path
