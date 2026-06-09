"""Testes dos servicos de insights e PowerPoint."""

import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from pptx import Presentation
from sqlalchemy import create_engine

from app.services.insights import InsightsService
from app.services.powerpoint import PowerPointService
from app.services.query_results import QueryResultStore


class ServiceTests(unittest.TestCase):
    """Valida contratos de consulta e apresentacao sem agentes."""

    def test_insights_service_runs_valid_query(self) -> None:
        """Uma consulta valida deve produzir resultado e referencia persistida."""
        engine = create_engine("sqlite:///:memory:")
        pd.DataFrame(
            {
                "QtdItens": [2, 3],
                "ValorUnitario": [10.0, 5.0],
            }
        ).to_sql("fact", engine, index=False)
        result_store = QueryResultStore()
        service = InsightsService(engine, result_store)

        result = service.run_sql_tool(
            "SELECT SUM(fact.QtdItens * fact.ValorUnitario) "
            "AS TotalVendas FROM fact"
        )

        self.assertTrue(result["ok"])
        self.assertTrue(result["result_id"])
        self.assertEqual(result["row_count"], 1)
        self.assertEqual(result["preview_rows"][0]["TotalVendas"], 35.0)

    def test_powerpoint_service_creates_multiple_table_slides(self) -> None:
        """Cada resultado deve originar um slide tabular no mesmo arquivo."""
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "slide.pptx"
            result_store = QueryResultStore()
            service = PowerPointService(output_path, result_store)
            with patch("app.services.powerpoint.datetime") as mock_datetime:
                mock_datetime.now.return_value = datetime(
                    2026, 6, 7, 2, 54, 13
                )
                service.start_report()
                first_result_id = result_store.save(
                    pd.DataFrame(
                        {
                            "UF": ["SP", "RJ"],
                            "TotalVendas": [100.0, 70.0],
                        }
                    )
                )
                second_result_id = result_store.save(
                    pd.DataFrame(
                        {
                            "Ano": [2024, 2025],
                            "TotalVendas": [120.0, 180.0],
                        }
                    )
                )

                first = service.create_table_slide(
                    result_id=first_result_id,
                    titulo="Vendas por UF",
                    headline="SP concentra o maior valor de vendas.",
                )
                second = service.create_table_slide(
                    result_id=second_result_id,
                    titulo="Evolucao anual",
                    headline="As vendas cresceram em 2025.",
                )

            self.assertTrue(first["ok"])
            self.assertEqual(second["slide_number"], 2)
            self.assertEqual(
                Path(first["pptx_path"]).name,
                "slide_20260607_025413.pptx",
            )
            self.assertTrue(Path(first["pptx_path"]).exists())
            self.assertGreater(Path(first["pptx_path"]).stat().st_size, 0)
            presentation = Presentation(first["pptx_path"])
            self.assertEqual(len(presentation.slides), 2)
            self.assertEqual(len(presentation.slides[0].shapes), 3)
            self.assertTrue(presentation.slides[0].shapes[2].has_table)
            self.assertFalse(presentation.slides[0].shapes[2].has_chart)


if __name__ == "__main__":
    unittest.main()
